import base64
from concurrent.futures import ThreadPoolExecutor
import io
from pathlib import Path
import tempfile
import unittest
import wave
from docx import Document
from fastapi.testclient import TestClient
from app import app, analyze
from vendor import engine
from sonification import sonify

TEXT = ('Tydlighet, öppenhet och respekt skapar möjligheter. '
        'Kunskap och ansvar är viktigt. Det är inte resultatet utan processen som räknas. '
        'Dialog, samarbete och återkoppling hjälper eleverna. '
        'Nyfikenhet, självständighet och kreativitet öppnar dörrar. '
        'Kvalitet, tillit och omtanke skapar trygghet. '
        'Jag läste boken och jag tänkte på innehållet. ')

class AppTests(unittest.TestCase):
    def setUp(self):
        self.client = TestClient(app)

    def test_audio_and_events_match_unmodified_engine(self):
        for text in ('Tydlighet, öppenhet och respekt är viktigt.', TEXT * 8 + 'En mening. ' * 100, 'Jag åt och jag sov och jag gick och jag läste och jag skrev.'):
            with self.subTest(text=text[:20]), tempfile.TemporaryDirectory() as directory:
                a, b = Path(directory)/'a.wav', Path(directory)/'b.wav'
                expected = engine.sonify(text, 'sv', a)
                actual, events = sonify(text, b)
                self.assertEqual(a.read_bytes(), b.read_bytes())
                self.assertEqual(expected['events'], len(events))
                self.assertEqual(expected['word_clock_event_count'], sum(e['kind']=='tick' for e in events))
                self.assertEqual(expected['audible_triad_count']*3, sum(e['kind']=='triad' for e in events))
                self.assertEqual(expected['density_alert_count'], sum(e['kind']=='density' for e in events))
                self.assertEqual(expected['extreme_wow_count'], sum(e['kind']=='wow' for e in events))
                self.assertTrue(all(0<=e['time']<actual['seconds'] for e in events))
        result=analyze(TEXT)
        self.assertGreater(result['report']['audio']['density_alert_count'],0)
        self.assertGreater(result['report']['audio']['extreme_wow_count'],0)

    def test_original_short_cluster_render_bug_is_fixed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'original.wav'
            with self.assertRaises(ValueError):
                engine.sonify(TEXT, 'sv', path)
        result = analyze(TEXT)
        self.assertEqual(result['report']['features'], engine.build_report(TEXT, 'sv')['features'])
        self.assertGreater(result['report']['audio']['seconds'], result['report']['audio']['timeline_seconds'])

    def test_text_and_download_payload(self):
        r=self.client.post('/api/analyze', data={'text':TEXT})
        self.assertEqual(r.status_code,200)
        result=r.json()
        self.assertEqual(result['report']['features'],engine.build_report(TEXT,'sv')['features'])
        self.assertEqual(r.headers['cache-control'],'no-store')
        self.assertIn('anomaly_index',result['csv'])
        with wave.open(io.BytesIO(base64.b64decode(result['audio_base64']))) as wav:
            self.assertEqual(wav.getframerate(),44100)
            self.assertGreater(wav.getnframes(),0)

    def test_uploads(self):
        doc=Document();doc.add_paragraph(TEXT);table=doc.add_table(rows=1,cols=1);table.cell(0,0).text='Extra text i tabell.'
        stream=io.BytesIO();doc.save(stream)
        for name,content in [('test.txt',TEXT.encode('utf-8')),('test.docx',stream.getvalue())]:
            r=self.client.post('/api/analyze',files={'file':(name,content)})
            self.assertEqual(r.status_code,200,r.text[:100])
            self.assertGreater(r.json()['report']['features']['words'],0)
        r=self.client.post('/api/analyze',files={'file':('bad.docx',b'not zip')})
        self.assertEqual(r.status_code,422)

    def test_validation(self):
        for text in ('','  ','!!!'):
            self.assertEqual(self.client.post('/api/analyze',data={'text':text}).status_code,422)
        self.assertEqual(self.client.post('/api/analyze',data={'text':'ord '*12001}).status_code,413)
        self.assertEqual(self.client.post('/api/analyze',files={'file':('x.pdf',b'pdf')}).status_code,422)
        self.assertEqual(self.client.post('/api/analyze',files={'file':('x.txt',b'\xff')}).status_code,422)
        self.assertEqual(self.client.post('/api/analyze',files={'file':('x.txt',b'x'*(2*1024*1024+1))}).status_code,413)
        self.assertEqual(self.client.post('/api/analyze',content=b'x'*(3*1024*1024)).status_code,413)

    def test_parallel_isolation(self):
        with ThreadPoolExecutor(max_workers=2) as pool:
            results=list(pool.map(analyze,[TEXT,'Ett enda ord.']))
        for text,result in zip([TEXT,'Ett enda ord.'],results):
            self.assertEqual(result['report']['features']['words'],len(engine.tokenize(text)))
            self.assertEqual(result['report']['audio']['events'],len(result['report']['timeline']))

    def test_static_site(self):
        self.assertEqual(self.client.get('/').status_code,200)
        self.assertEqual(self.client.get('/api/health').json()['engine'],'2.10')
        self.assertEqual(self.client.get('/static/soundbank/00_ordklocka_tick_var_30e_ord.wav').status_code,200)

if __name__=='__main__': unittest.main()
